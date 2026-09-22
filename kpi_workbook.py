"""Builds output/Core_AC_DC_KPI.xlsx: the KPI fact table, its dimensions and its parameters.

This workbook reports one row per BusinessDay x Timestamp x Country x Case x VoltageLevel.
The tool runs a single snapshot, so BusinessDay and Timestamp are one fixed pair per run and
the grain that actually varies is Country x Case x VoltageLevel.

It does not replace output/kpi_results.xlsx. That one carries the full per-contingency
series; this one carries N-0, all-COs-combined and per-CO-count only. Both are written by
every run.
"""

import datetime
import logging

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

from rosc_acdc import config, dim_kpi, kpis
from rosc_acdc.paths import output_path

logger = logging.getLogger(__name__)

CASE_BASE = "BaseCase"
CASE_CONTINGENCY = "Contingency"
CASE_ALL = "All"

KPI_DATA_SHEET = "KPI_Data"
PERF_SHEET = "Perf_Computation"
PARAMETERS_SHEET = "Parameters"

# Every tab the workbook carries, in order. Four are filled; the rest exist by name only,
# which is all this build needs them to do.
EMPTY_SHEETS = [
    "README", "Raw_Sample", "Dim_Date", "Dim_Time", "Dim_Country", "Dim_Case",
    "Dim_VoltageLevel",
]
SHEET_ORDER = [
    "README", KPI_DATA_SHEET, PERF_SHEET, "Raw_Sample", "Dim_Date", "Dim_Time",
    "Dim_Country", "Dim_Case", "Dim_VoltageLevel", dim_kpi.SHEET_NAME, PARAMETERS_SHEET,
]

KEY_COLUMNS = [
    "BusinessDay", "Timestamp", "TimestampID", "Country", "Case", "VoltageLevel_kV",
]
FLOW_COLUMNS = ["MAE_A", "RMSE_A", "P95_Err_A", "Max_Err_A"]
LOADING_COLUMNS = [
    "Mean_LoadingDev_pp", "Mean_Signed_Margin_Err_A", "Sum_NearLimit_LoadingDev_pp",
    "N_NearLimit_Obs", "NearLimit_Mean_LoadingDev_pp", "N_NearLimit_Underest",
]
VIOLATION_COLUMNS = ["N_FN", "N_FP", "Missed_Overload_Rate", "False_Alarm_Rate"]
SEVERITY_COLUMNS = [
    "Missed_Overload_Volume_A", "False_Overload_Volume_A", "Avg_FN_Loading_AC_pct",
    "Avg_FP_Loading_DC_pct",
]
RANKING_COLUMNS = ["TopN_Overlap_Count", "TopN_N", "TopN_Overlap_Rate", "WorstCase_Match_Rate"]
PRIORITY2_COLUMNS = [
    "NearLimit_Underest_Rate", "Missed_Critical_CO_Rate", "False_Critical_CO_Rate",
    "Mean_WorstCO_SevDev_A",
]
PRIMITIVE_COLUMNS = [
    "N_Obs", "N_CO", "N_Elements", "Sum_AbsErr_A", "Sum_SqErr_A", "Sum_SignedErr_A",
    "Max_AbsErr_A", "Sum_LoadingDev_pp", "Sum_MarginDev_A", "N_AC_Viol", "N_DC_Viol",
    "N_TP", "N_TN", "Sum_MissedOverload_A", "Sum_FalseOverload_A", "Sum_FN_Loading_AC_pct",
    "Sum_FP_Loading_DC_pct",
]

KPI_DATA_COLUMNS = (
    KEY_COLUMNS + FLOW_COLUMNS + LOADING_COLUMNS + VIOLATION_COLUMNS + SEVERITY_COLUMNS
    + RANKING_COLUMNS + PRIORITY2_COLUMNS + PRIMITIVE_COLUMNS
)

PERF_COLUMNS = [
    "BusinessDay", "Timestamp", "TimestampID", "Process", "Process_Type", "Model",
    "Computation_Time_s", "N_Elements_Evaluated", "N_Contingencies",
]


# ---------------------------------------------------------------------------
# Run identity
# ---------------------------------------------------------------------------


def run_timestamp(network) -> tuple:
    """(BusinessDay, Timestamp, TimestampID) for this run, from the network's own case date.

    `case_date` is the scenario time the CGMES import read out of the model's FullModel
    header, which is the time the network actually represents. The CGM filename encodes a
    time too and is not used: the two disagree on the Belgovia sample by about 25 hours,
    and the .xiidm import path never sees a CGMES filename at all.

    BusinessDay is the UTC date, with no timezone conversion. TimestampID is the UTC hour,
    so it names the hour-slot this snapshot represents rather than counting runs.

    Excel holds no timezone, so the returned Timestamp is naive UTC.
    """
    case_date = network.case_date
    if case_date.tzinfo is not None:
        case_date = case_date.astimezone(datetime.timezone.utc).replace(tzinfo=None)
    return case_date.date(), case_date, int(case_date.hour)


# ---------------------------------------------------------------------------
# Observations
# ---------------------------------------------------------------------------


def _observations(comparison, element_ids, contingency_ids, case, locations):
    """One comparison dataset as observation rows carrying their grouping keys."""
    observations = comparison.copy()
    observations["Element_Id"] = element_ids.reindex(comparison.index)
    observations["Contingency_Id"] = (
        contingency_ids.reindex(comparison.index) if contingency_ids is not None
        else pd.Series(np.nan, index=comparison.index)
    )
    observations["Case"] = case

    located = locations.reindex(observations["Element_Id"])
    observations["Country"] = located["Country"].to_numpy()
    observations["VoltageLevel_kV"] = located["VoltageLevel_kV"].to_numpy()
    return observations.reset_index(drop=True)


def build_observations(locations, base_case_comparison=None, base_case_ids=None,
                        sa_comparison=None, sa_element_ids=None, sa_contingency_ids=None):
    """Pool the base-case and contingency observations into one table with grouping keys.

    The Case populations are then slices of this table, which is what makes `All` the union
    of the other two rather than a third calculation: the same KPI functions run over the
    same observations, just a wider selection of them. It is also the only way to get a
    correct pooled P95, which cannot be rebuilt from summed primitives.
    """
    frames = []

    if base_case_comparison is not None and not base_case_comparison.empty:
        frames.append(_observations(
            base_case_comparison, base_case_ids, None, CASE_BASE, locations,
        ))

    if sa_comparison is not None and not sa_comparison.empty:
        # The SA dataset carries the intact network as rows with no contingency id
        # (security_analysis.build_shortlist_and_full_comparison concatenates them in).
        # They are the base case a second time, so counting them here would double-count
        # every N-0 observation in the All rows. They are dropped explicitly rather than
        # left to prepare_comparison's missing-DC-value filter, which happens to remove
        # them today only because of the unrelated data-sourcing bug in
        # PRIORITY1_KPI_CHANGES.md 5.3 - fixing that bug must not silently reintroduce
        # the double count here.
        contingency_only = sa_contingency_ids.reindex(sa_comparison.index).notna()
        dropped = int((~contingency_only).sum())
        if dropped:
            logger.info(
                "KPI_Data: excluded %d intact-network row(s) from the Contingency case; "
                "the base case is already counted under Case=BaseCase", dropped,
            )
        frames.append(_observations(
            sa_comparison[contingency_only], sa_element_ids, sa_contingency_ids,
            CASE_CONTINGENCY, locations,
        ))

    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True)


# ---------------------------------------------------------------------------
# One KPI_Data row
# ---------------------------------------------------------------------------


def _kpi_row(observations: pd.DataFrame) -> dict:
    """Every KPI and primitive for one (Country, Case, VoltageLevel) slice."""
    element_ids = observations["Element_Id"]
    contingency_ids = observations["Contingency_Id"]

    flow = kpis.flow_deviation_kpis(observations)
    loading = kpis.loading_limit_kpis(observations)
    near_limit = kpis.near_limit_kpis(observations)
    violation = kpis.violation_kpis(observations)
    severity = kpis.severity_kpis(observations, violation)
    ranking = kpis.element_ranking_kpis(observations, element_ids)
    co_level = kpis.co_level_kpis(observations, contingency_ids)
    primitives = kpis.raw_primitives(observations, violation, element_ids, contingency_ids)

    near_limit_label = (
        f"Near-Limit ({config.KPI_NEAR_LIMIT_THRESHOLD_PCT}%+) Mean Loading % Deviation"
    )

    row = {
        "MAE_A": flow["MAE (A)"],
        "RMSE_A": flow["RMSE (A)"],
        "P95_Err_A": flow["P95 Error (A)"],
        "Max_Err_A": flow["Max Error (A)"],

        "Mean_LoadingDev_pp": loading["Mean Loading % Deviation (DC-AC)"],
        "Mean_Signed_Margin_Err_A": loading["Mean Margin Error (DC-AC) (A)"],
        "Sum_NearLimit_LoadingDev_pp": near_limit["Sum_NearLimit_LoadingDev_pp"],
        "N_NearLimit_Obs": near_limit["N_NearLimit_Obs"],
        "NearLimit_Mean_LoadingDev_pp": loading[near_limit_label],
        "N_NearLimit_Underest": near_limit["N_NearLimit_Underest"],

        "N_FN": violation["False Negatives"],
        "N_FP": violation["False Positives"],
        "Missed_Overload_Rate": violation["Missed Overload Rate"],
        "False_Alarm_Rate": violation["False Alarm Rate"],

        "Missed_Overload_Volume_A": severity["Missed Overload Volume (A)"],
        "False_Overload_Volume_A": severity["False Overload Volume (A)"],
        "Avg_FN_Loading_AC_pct": severity["False Negatives Avg Loading %"],
        "Avg_FP_Loading_DC_pct": severity["False Positives Avg Loading %"],

        "TopN_Overlap_Count": ranking["TopN_Overlap_Count"],
        "TopN_N": ranking["TopN_N"],
        "TopN_Overlap_Rate": ranking["TopN_Overlap_Rate"],
        "WorstCase_Match_Rate": ranking["WorstCase_Match_Rate"],

        "NearLimit_Underest_Rate": near_limit["NearLimit_Underest_Rate"],
        "Missed_Critical_CO_Rate": co_level["Missed_Critical_CO_Rate"],
        "False_Critical_CO_Rate": co_level["False_Critical_CO_Rate"],
        "Mean_WorstCO_SevDev_A": co_level["Mean_WorstCO_SevDev_A"],
    }
    row.update(primitives)
    return row


def build_kpi_data(observations: pd.DataFrame, business_day, timestamp, timestamp_id):
    """The KPI_Data fact table: one row per Country x Case x VoltageLevel for this snapshot.

    `All` pools the BaseCase and Contingency observations and runs the same KPI functions
    over the union. It is never an arithmetic combination of the other two rows - that
    would get P95 wrong, and would re-derive every ratio from ratios.
    """
    if observations.empty:
        logger.warning("KPI_Data: no observations, writing an empty sheet")
        return pd.DataFrame(columns=KPI_DATA_COLUMNS)

    rows = []
    group_keys = ["Country", "VoltageLevel_kV"]
    for (country, voltage_level), located in observations.groupby(group_keys, dropna=False):
        populations = [
            (CASE_BASE, located[located["Case"] == CASE_BASE]),
            (CASE_CONTINGENCY, located[located["Case"] == CASE_CONTINGENCY]),
            (CASE_ALL, located),
        ]
        for case, population in populations:
            if population.empty:
                continue
            row = {
                "BusinessDay": business_day,
                "Timestamp": timestamp,
                "TimestampID": timestamp_id,
                "Country": country,
                "Case": case,
                "VoltageLevel_kV": voltage_level,
            }
            row.update(_kpi_row(population))
            rows.append(row)

    kpi_data = pd.DataFrame(rows, columns=KPI_DATA_COLUMNS)
    logger.info(
        "KPI_Data: %d rows over %d country/voltage-level combination(s)",
        len(kpi_data), kpi_data.groupby(group_keys, dropna=False).ngroups if len(kpi_data) else 0,
    )
    return kpi_data


# ---------------------------------------------------------------------------
# Perf_Computation
# ---------------------------------------------------------------------------


def build_perf_computation(business_day, timestamp, timestamp_id, stage_times,
                            n_elements_evaluated, n_contingencies):
    """One row per computation stage: the 9 confirmed columns, and nothing else.

    No AC/DC totals and no speed-up ratio. Those are not among Perf_Computation's confirmed
    columns, and this workbook writes only what the schema carries. kpi_results.xlsx still
    reports both, so nothing is lost - see PRIORITY1_KPI_CHANGES.md 4.8.

    `stage_times` is an ordered (Process, Process_Type, Model, seconds) sequence; a stage
    that did not run is left out by the caller rather than written as zero.
    """
    rows = [
        {
            "BusinessDay": business_day,
            "Timestamp": timestamp,
            "TimestampID": timestamp_id,
            "Process": process,
            "Process_Type": process_type,
            "Model": model,
            "Computation_Time_s": seconds,
            "N_Elements_Evaluated": n_elements_evaluated,
            "N_Contingencies": n_contingencies,
        }
        for process, process_type, model, seconds in stage_times
        if seconds is not None
    ]
    return pd.DataFrame(rows, columns=PERF_COLUMNS)


# ---------------------------------------------------------------------------
# Parameters
# ---------------------------------------------------------------------------


def build_parameters() -> pd.DataFrame:
    """The config this run actually used, reported after the fact.

    Nobody edits this sheet to change a run; the input is rosc_acdc/config.py plus the
    git-ignored input/config_local.py override.
    """
    rows = [
        ("Violation threshold (% loading)", config.KPI_VIOLATION_THRESHOLD_PCT,
         "An element is in violation at or above this loading (>=). Drives N_AC_Viol, "
         "N_DC_Viol, N_FN, N_FP, N_TP, N_TN."),
        ("Near-limit threshold (% loading, AC)", config.KPI_NEAR_LIMIT_THRESHOLD_PCT,
         "Near-limit KPIs look only at elements whose AC loading reaches this. Distinct "
         "from the violation threshold above."),
        ("Top-N for critical element overlap", config.KPI_TOP_N,
         "N for TopN_Overlap_Count and the TopN_N column."),
        ("Business days simulated", config.KPI_BUSINESS_DAYS_SIMULATED,
         "Report-only: this build runs a single snapshot."),
        ("Timestamps per business day", config.KPI_TIMESTAMPS_PER_BUSINESS_DAY,
         "Report-only: this build runs a single snapshot."),
        ("Contingencies per country / voltage level",
         config.KPI_CONTINGENCIES_PER_COUNTRY_VOLTAGE_LEVEL,
         "Report-only. Perf_Computation's N_Contingencies reports the contingencies the "
         "security analysis actually ran, not this value."),
        ("Flow unit", config.KPI_FLOW_UNIT, "Display and labelling only."),
        ("Default country (unresolved)", config.KPI_FALLBACK_COUNTRY,
         "Reported for elements whose substation carries no country, which is every "
         "element in a model imported without its boundary files."),
    ]
    return pd.DataFrame(rows, columns=["Parameter", "Value", "Comment"])


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------


def _default_format(column: str):
    """The format for a column Dim_KPI does not govern - a primitive or a key column.

    UNCONFIRMED: counts as #,##0 and sums as #,##0.0 is this build's suggestion, per
    04_business_rules_and_kpi_catalog.md's own note that it awaits Mahir. Assigned by name
    rather than from a list so that adding a primitive column cannot silently leave it
    unformatted.
    """
    if column == "BusinessDay":
        return "yyyy-mm-dd"
    if column == "Timestamp":
        return "yyyy-mm-dd hh:mm"
    if column.startswith(("Sum_", "Max_")):
        return dim_kpi.PRIMITIVE_SUM_FORMAT
    if column.startswith(("N_", "TopN_")) or column in ("TimestampID", "VoltageLevel_kV"):
        return dim_kpi.PRIMITIVE_COUNT_FORMAT
    return None


def _column_styles(catalog: pd.DataFrame, columns) -> dict:
    """column -> {format, direction, target}: Dim_KPI for the 25 KPIs, defaults elsewhere."""
    styles = {
        column: {"format": _default_format(column), "direction": None, "target": None}
        for column in columns
    }

    # Dim_KPI wins wherever it governs a column: it is the single source for the 25 KPIs,
    # for the number format and for the Direction/Target that drive the pass/fail flag.
    styles.update(dim_kpi.column_attributes(catalog, columns))
    return styles


# A KPI value that misses its Dim_KPI target. Fill and font rather than a conditional
# formatting rule, so the flag is baked into the file and survives being read back.
FAIL_FILL = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
FAIL_FONT = Font(color="FF9C0006", bold=True)


def _apply_styles(worksheet, columns, styles) -> int:
    """Format every written column and flag the values that miss their target.

    Returns how many cells were flagged. Joined by column name rather than by position:
    the workbook is built to field names, not to Excel column letters.
    """
    flagged = 0
    for index, column in enumerate(columns, start=1):
        style = styles.get(column)
        if not style:
            continue
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=index)
            if style["format"]:
                cell.number_format = style["format"]
            if dim_kpi.fails_target(cell.value, style["direction"], style["target"]):
                cell.fill = FAIL_FILL
                cell.font = FAIL_FONT
                flagged += 1
    return flagged


def write_workbook(kpi_data, perf_computation, parameters, catalog, filename=None) -> str:
    """Write every tab of Core_AC_DC_KPI.xlsx and return the path.

    All 11 tabs are created in the workbook's own order; the 7 that are placeholders exist
    by name with no content. The file is overwritten on every run - no versioning, no
    timestamp in the name.
    """
    if filename is None:
        filename = config.KPI_WORKBOOK_FILENAME
    path = output_path(filename)
    styles = _column_styles(catalog, list(kpi_data.columns) + list(perf_computation.columns))

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        kpi_data.to_excel(writer, sheet_name=KPI_DATA_SHEET, index=False)
        perf_computation.to_excel(writer, sheet_name=PERF_SHEET, index=False)
        catalog.to_excel(writer, sheet_name=dim_kpi.SHEET_NAME, index=False)
        parameters.to_excel(writer, sheet_name=PARAMETERS_SHEET, index=False)

        book = writer.book
        for sheet in EMPTY_SHEETS:
            if sheet not in book.sheetnames:
                book.create_sheet(sheet)

        flagged = _apply_styles(writer.sheets[KPI_DATA_SHEET], list(kpi_data.columns), styles)
        flagged += _apply_styles(
            writer.sheets[PERF_SHEET], list(perf_computation.columns), styles,
        )

        book._sheets = [book[name] for name in SHEET_ORDER if name in book.sheetnames]

    logger.info(
        "Wrote %s: %d KPI_Data row(s), %d Perf_Computation row(s), %d Dim_KPI row(s), "
        "%d value(s) flagged against their Dim_KPI target",
        path, len(kpi_data), len(perf_computation), len(catalog), flagged,
    )
    return path


def build_and_write(network, locations, stage_times, n_elements_evaluated, n_contingencies,
                     base_case_comparison=None, base_case_ids=None, sa_comparison=None,
                     sa_element_ids=None, sa_contingency_ids=None) -> str:
    """Build every tab from one run's results and write the workbook."""
    business_day, timestamp, timestamp_id = run_timestamp(network)
    logger.info(
        "KPI workbook run identity: BusinessDay=%s Timestamp=%s TimestampID=%s",
        business_day, timestamp, timestamp_id,
    )

    observations = build_observations(
        locations, base_case_comparison, base_case_ids,
        sa_comparison, sa_element_ids, sa_contingency_ids,
    )
    kpi_data = build_kpi_data(observations, business_day, timestamp, timestamp_id)

    # K19/K20 are not columns: they are K01/K02 and K12 read on the BaseCase rows. Logged
    # rather than written, since KPI_Data's schema is fixed and has no place for them.
    readings = dim_kpi.base_case_readings(kpi_data, CASE_BASE)
    logger.info(
        "Dim_KPI K19 (N-0 MAE/RMSE) = %.3f / %.3f A; K20 (N-0 Missed Overload Volume) = %.3f A",
        readings["K19_MAE_A"], readings["K19_RMSE_A"],
        readings["K20_Missed_Overload_Volume_A"],
    )

    perf_computation = build_perf_computation(
        business_day, timestamp, timestamp_id, stage_times, n_elements_evaluated,
        n_contingencies,
    )
    return write_workbook(kpi_data, perf_computation, build_parameters(), dim_kpi.load_catalog())
